import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree

from src.template_parser import TemplateParser, _W_NS


class TemplateParserTableBordersTests(unittest.TestCase):
    def test_parse_table_borders_patterns(self) -> None:
        document = Document()
        for _ in range(5):
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "x"

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "tables.docx"
            document.save(str(path))
            xml_bytes = _update_table_borders(path)
            _rewrite_document_xml(path, xml_bytes)

            parser = TemplateParser()
            result = parser.parse(str(path))

        table_borders = result.meta.get("table_borders", {})
        summary = table_borders.get("summary", {})
        self.assertEqual(summary.get("grid"), 1)
        self.assertEqual(summary.get("outer_only"), 1)
        self.assertEqual(summary.get("inner_only"), 1)
        self.assertEqual(summary.get("none"), 1)
        self.assertEqual(summary.get("mixed"), 1)
        tables = table_borders.get("tables", [])
        self.assertEqual(len(tables), 5)
        patterns = [table.get("pattern") for table in tables]
        self.assertEqual(
            patterns,
            ["grid", "outer_only", "inner_only", "none", "mixed"],
        )


def _update_table_borders(path: Path) -> bytes:
    with ZipFile(path) as archive:
        document_bytes = archive.read("word/document.xml")
    root = etree.fromstring(document_bytes)
    tables = root.findall(".//w:tbl", namespaces={"w": _W_NS})
    patterns = [
        ("grid", {"top", "bottom", "left", "right", "insideH", "insideV"}),
        ("outer_only", {"top", "bottom", "left", "right"}),
        ("inner_only", {"insideH", "insideV"}),
        ("none", set()),
        ("mixed", {"top", "left"}),
    ]
    for table, (_, borders) in zip(tables, patterns, strict=False):
        _apply_tbl_borders(table, borders)
    return etree.tostring(root, xml_declaration=True, encoding="utf-8")


def _apply_tbl_borders(table: etree._Element, borders: set[str]) -> None:
    tbl_pr = table.find(f"{{{_W_NS}}}tblPr")
    if tbl_pr is None:
        tbl_pr = etree.SubElement(table, f"{{{_W_NS}}}tblPr")
    tbl_borders = tbl_pr.find(f"{{{_W_NS}}}tblBorders")
    if tbl_borders is None:
        tbl_borders = etree.SubElement(tbl_pr, f"{{{_W_NS}}}tblBorders")
    for key in ("top", "bottom", "left", "right", "insideH", "insideV"):
        child = tbl_borders.find(f"{{{_W_NS}}}{key}")
        if child is None:
            child = etree.SubElement(tbl_borders, f"{{{_W_NS}}}{key}")
        value = "single" if key in borders else "nil"
        child.set(f"{{{_W_NS}}}val", value)


def _rewrite_document_xml(path: Path, xml_bytes: bytes) -> None:
    with ZipFile(path) as archive:
        data = {name: archive.read(name) for name in archive.namelist()}
    data["word/document.xml"] = xml_bytes
    with ZipFile(path, "w") as archive:
        for name, content in data.items():
            archive.writestr(name, content)


if __name__ == "__main__":
    unittest.main()
